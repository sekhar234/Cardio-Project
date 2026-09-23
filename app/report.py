"""Downloadable research report (Markdown and Word), built deterministically from the database.

No LLM is involved at this stage: the report is a rendering of verified claims, the validated
briefing, coverage gaps, conflicts and the source/crawl audit trail.
"""
from __future__ import annotations

import io
import re
from typing import Any

from sqlalchemy import select

from . import db
from .taxonomy import DIMENSIONS

VERDICT_LABEL = {"supported": "Supported", "partially_supported": "Partially supported",
                 "unsupported": "Rejected"}


async def collect(run_id: str) -> dict[str, Any] | None:
    async with db.session() as s:
        run = await s.get(db.Run, run_id)
        if not run:
            return None
        claims = list((await s.execute(select(db.Claim).where(db.Claim.run_id == run_id)
                                       .order_by(db.Claim.seq))).scalars())
        sources = {x.id: x for x in (await s.execute(select(db.Source).where(db.Source.run_id == run_id))).scalars()}
        conflicts = list((await s.execute(select(db.Conflict).where(db.Conflict.run_id == run_id))).scalars())
    return {"run": run, "claims": claims, "sources": sources, "conflicts": conflicts}


def _cites(ids: list[int]) -> str:
    return " " + " ".join(f"[C{i}]" for i in ids)


def build_blocks(data: dict[str, Any]) -> list[tuple[str, Any]]:
    """A tiny document model: (kind, payload) with kinds h1/h2/h3/p/bullets/table/note."""
    run, claims, sources, conflicts = data["run"], data["claims"], data["sources"], data["conflicts"]
    verified = [c for c in claims if c.verdict in ("supported", "partially_supported")]
    brief = run.brief or {}
    stats = run.stats or {}
    b: list[tuple[str, Any]] = []
    b.append(("h1", f"City intelligence brief: {run.city}, {run.country}"))
    b.append(("p", f"Generated {run.finished_at:%d %b %Y %H:%M} UTC · run {run.id} · "
                   f"{stats.get('sources_fetched', 0)} sources read, {stats.get('sources_blocked', 0)} not crawled "
                   f"(permissions), {stats.get('claims_supported', 0) + stats.get('claims_partial', 0)} verified "
                   f"claims, {stats.get('claims_rejected', 0)} rejected by the fact checker."
              if run.finished_at else f"Run {run.id} ({run.status})"))
    b.append(("note", "How to read this: every statement cites verified claims [C#] listed in the Evidence "
                      "appendix with the exact source quote. Items marked NATIONAL/REGIONAL are not city-specific. "
                      "Opportunities and risks are analysis (inference), not facts, and state their assumptions."))

    if brief.get("executive_summary"):
        b.append(("h2", "Executive summary"))
        b.append(("bullets", [p["text"] + _cites(p["claims"]) for p in brief["executive_summary"]]))

    # Key statistics table straight from claims (no LLM wording)
    stats_rows = [c for c in verified if c.claim_type == "statistic"]
    if stats_rows:
        b.append(("h2", "Key figures (as stated in sources)"))
        b.append(("table", [["Metric", "Value", "Year", "Geography", "Ref"]] + [
            [(c.metric_key or "other").replace("_", " "),
             f"{c.value:g} {c.unit}".strip() if c.value is not None else "—", c.year or "n/s",
             f"{c.geography_level}{' ⚠ not city' if c.not_city_level else ''}", f"C{c.seq}"]
            for c in stats_rows[:30]]))

    coverage = run.coverage or {}
    for d in DIMENSIONS:
        b.append(("h2", d.title))
        cov = coverage.get(d.key, {})
        if cov:
            b.append(("p", f"Evidence status: {cov.get('status', 'n/a').upper()} — {cov.get('verified_claims', 0)} "
                           f"verified claims ({cov.get('city_level_claims', 0)} city-level)."))
        pts = (brief.get("sections") or {}).get(d.key, {}).get("points", [])
        if pts:
            b.append(("bullets", [p["text"] + _cites(p["claims"]) for p in pts]))
        missing = [q["question"] for q in cov.get("questions", []) if q["status"] != "answered"]
        if missing:
            b.append(("h3", "Not established by the research"))
            b.append(("bullets", [f"{q} ({'partial' if any(x['question'] == q and x['status'] == 'partial' for x in cov['questions']) else 'not found'})"
                                  for q in missing]))

    for key, title in (("opportunities", "Opportunities (analysis)"), ("risks", "Risks (analysis)")):
        items = brief.get(key) or []
        if items:
            b.append(("h2", title))
            b.append(("bullets", [f"{p['text']}{_cites(p['claims'])} — Assumption: {p.get('assumption', 'n/s')}"
                                  for p in items]))

    if brief.get("meeting_questions"):
        b.append(("h2", "Questions to raise in the meeting"))
        b.append(("bullets", brief["meeting_questions"]))

    if conflicts:
        b.append(("h2", "Conflicting and time-varying figures"))
        b.append(("bullets", [c.description for c in conflicts]))

    b.append(("h2", "Evidence appendix"))
    for c in verified:
        src = sources.get(c.source_id)
        flag = " ⚠ NATIONAL/REGIONAL DATA" if c.not_city_level else ""
        b.append(("p", f"C{c.seq} [{VERDICT_LABEL[c.verdict]}{flag}] {c.statement}"))
        b.append(("note", f"“{c.quote}” — {src.title or src.domain}, {src.url} "
                          f"(published {src.published_date or 'n/s'}; retrieved "
                          f"{src.fetched_at:%Y-%m-%d} ) · Check: {c.verdict_reason}" if src and src.fetched_at
                  else f"“{c.quote}”"))

    rejected = [c for c in claims if c.verdict == "unsupported"]
    b.append(("h2", "Method and audit trail"))
    b.append(("p", "Pipeline: plan → web search → crawlability check (robots.txt + terms policy, before any "
                   "page request) → fetch → claim extraction with verbatim quotes → independent fact check "
                   "(deterministic quote & number check, then a separate model) → coverage assessment with "
                   "targeted follow-up research → conflict detection → knowledge graph (Graphiti) → briefing "
                   "with citation and number guard."))
    b.append(("p", f"Claims rejected by the fact checker: {len(rejected)} (kept in the audit log, excluded "
                   f"from this report). Generated statements dropped by the citation guard: "
                   f"{len(brief.get('dropped', []))}."))
    b.append(("table", [["Source", "Tier", "Crawl decision", "Status"]] + [
        [f"{x.domain} — {(x.title or '')[:60]}", str(x.credibility_tier),
         ("allowed" if x.crawl_allowed else "blocked") + f": {x.crawl_reason[:80]}", x.status]
        for x in sorted(sources.values(), key=lambda x: (x.crawl_allowed is not True, x.credibility_tier))]))
    return b


def to_markdown(blocks: list[tuple[str, Any]]) -> str:
    out = []
    for kind, val in blocks:
        if kind == "h1":
            out.append(f"# {val}\n")
        elif kind == "h2":
            out.append(f"\n## {val}\n")
        elif kind == "h3":
            out.append(f"\n**{val}**\n")
        elif kind == "p":
            out.append(f"{val}\n")
        elif kind == "note":
            out.append(f"> {val}\n")
        elif kind == "bullets":
            out.extend(f"- {x}" for x in val)
            out.append("")
        elif kind == "table":
            head, *rows = val
            out.append("| " + " | ".join(head) + " |")
            out.append("|" + "---|" * len(head))
            out.extend("| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in rows)
            out.append("")
    return "\n".join(out)


def to_docx(blocks: list[tuple[str, Any]]) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for kind, val in blocks:
        if kind == "h1":
            doc.add_heading(val, level=0)
        elif kind == "h2":
            doc.add_heading(val, level=1)
        elif kind == "h3":
            doc.add_heading(val, level=2)
        elif kind == "p":
            doc.add_paragraph(val)
        elif kind == "note":
            p = doc.add_paragraph()
            r = p.add_run(val)
            r.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif kind == "bullets":
            for x in val:
                doc.add_paragraph(x, style="List Bullet")
        elif kind == "table":
            head, *rows = val
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Light Grid Accent 1"
            for i, h in enumerate(head):
                t.rows[0].cells[i].text = h
            for r in rows:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = str(v)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def filename(run: db.Run, ext: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", f"C4C_{run.city}_{run.country}").strip("_") + f"_brief.{ext}"
